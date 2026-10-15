#!/usr/bin/env python3
"""Fiscalización externa — observaciones de la ASECAM por municipio.

    backend/datalab/venv/bin/python backend/datalab/pipeline/build_auditoria.py --descargar
    backend/datalab/venv/bin/python backend/datalab/pipeline/build_auditoria.py
    backend/datalab/venv/bin/python backend/datalab/pipeline/build_auditoria.py --autocomprobar

Noveno pipeline. Salida: frontend/data/electoral/auditoria.json

QUÉ ES Y QUÉ NO ES
──────────────────
Son observaciones de la Auditoría Superior del Estado de Campeche sobre las
cuentas públicas de 2023 y 2024. Es FISCALIZACIÓN EXTERNA sobre el ejercicio
de recursos públicos municipales — no es gasto de campaña, que es otra cosa
aunque las dos sean dinero. Por eso alimenta el grafo de riesgos y no el
módulo de finanzas.

Y como en forensia: una observación de auditoría NO es una sanción ni una
sentencia. Un pliego de observaciones es el inicio de un procedimiento, no su
resultado. Se publica el conteo, nunca una conclusión sobre él.

EL CERTIFICADO DE ASECAM ESTÁ INCOMPLETO, Y NO SE SILENCIA
──────────────────────────────────────────────────────────
asecam.gob.mx sirve su certificado sin el intermedio, así que cualquier cliente
estricto falla:

    SSLCertVerificationError: unable to get local issuer certificate  (código 20)

La salida fácil es desactivar la verificación (`verify=False`, `curl -k`). No se
hace: eso convierte una descarga de una fuente oficial en una descarga de
cualquiera que se ponga en medio.

Lo que se hace es COMPLETAR LA CADENA. El propio certificado declara dónde está
su intermedio, en la extensión Authority Information Access:

    CA Issuers - URI:http://cacerts.rapidssl.com/RapidSSLTLSRSACAG1.crt

Ese intermedio se descargó una vez, se guardó en el repositorio
(uploads/asecam/RapidSSLTLSRSACAG1.pem) y se añade al almacén de confianza del
sistema. La verificación sigue ACTIVA; lo único que cambia es que ahora puede
completarse. Un certificado es autoverificable —está firmado por la raíz de
DigiCert, que sí está en el sistema—, así que traerlo por HTTP no debilita nada.

VÍAS DESCARTADAS, CON SU RAZÓN
──────────────────────────────
  datos.gob.mx     Su API de CKAN devuelve 403 desde este entorno
                   (`/busca/api/3/action/package_search`). No es una elección:
                   no responde.

  ASF federal      asf.gob.mx publica 37 PDF en su sección de informes y cero
                   archivos procesables. Su portal de datos, asfdatos.gob.mx,
                   es un sistema de consulta web ("Sistema Público de Consulta
                   de Auditorías"), no un catálogo de descarga. Queda BLOQUEADO
                   por falta de fuente procesable, no por decisión.

  Archivos ASECAM  Los 10 archivos anteriores a 2023 tienen otro esquema (el de
  de años previos  2022 trae una sola hoja "Pliegos", con otras columnas). Se
                   dejan fuera en vez de forzar un mapeo entre formatos que no
                   se ha verificado.
"""

import argparse
import json
import pathlib
import re
import ssl
import sys
import unicodedata
import urllib.request

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl. Usa backend/datalab/venv/bin/python.")

RAIZ = pathlib.Path(__file__).resolve().parents[3]
CRUDO = RAIZ / "backend/datalab/uploads/asecam"
SALIDA = RAIZ / "frontend/data/electoral/auditoria.json"

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from build_electoral import CVE_MUN, NOMBRE_MUN, norm  # noqa: E402

FUENTE = "ASECAM — Auditoría Superior del Estado de Campeche, datos abiertos"
URL = ("https://asecam.gob.mx/pagina/Informes/2025/datos_abiertos/"
       "Datos_abiertos.xlsx")
ARCHIVO = "Datos_abiertos.xlsx"
INTERMEDIO = "RapidSSLTLSRSACAG1.pem"
FECHA_CORTE = "2025-11"

# Clasificación por TIPO de entidad, no por si el nombre menciona un municipio.
# La diferencia importa: "Instituto Tecnológico Superior de Champotón" es un
# organismo ESTATAL que está en Champotón, no un organismo del municipio.
# Clasificarlo como municipal le colgaría al municipio una observación que no
# es suya.
PATRONES = [
    ("ayuntamiento", r"^(?:H\.?\s+)?AYUNTAMIENTO\s+DE\s+(.+)$"),
    ("organismo_municipal", r"^SISTEMA\s+MUNICIPAL\s+DE\s+AGUA\s+POTABLE.*?\b(?:DE|DEL\s+MUNICIPIO\s+DE)\s+(.+)$"),
    ("organismo_municipal", r"^SISTEMA(?:\s+MUNICIPAL)?\s+PARA\s+EL\s+DESARROLLO\s+INTEGRAL\s+DE\s+LA\s+FAMILIA\s+(?:EN\s+EL\s+MUNICIPIO\s+DE|DEL\s+MUNICIPIO\s+DE|EN\s+EL\s+MUNICIPIO|EN\s+MUNICIPIO\s+DE|EN)\s+(.+)$"),
]

ENCABEZADO = ["Núm", "Entidad Fiscalizada", "Objeto de la Auditoría",
              "Observaciones", "Acciones"]


def contexto_tls():
    """Almacén del sistema MÁS el intermedio que ASECAM no envía.

    No se desactiva la verificación: se le da a Python la pieza que falta para
    poder verificar. Si el intermedio no está en el repositorio, el pipeline
    se detiene en vez de bajar por un canal sin verificar.
    """
    pem = CRUDO / INTERMEDIO
    if not pem.exists():
        raise SystemExit(
            f"Falta el certificado intermedio en {pem}.\n"
            f"ASECAM no lo envía en su cadena, así que hay que traerlo del "
            f"punto que declara su propio certificado:\n"
            f"  curl -s http://cacerts.rapidssl.com/RapidSSLTLSRSACAG1.crt "
            f"| openssl x509 -inform DER -out {pem}\n"
            f"NO se descarga con la verificación desactivada.")
    ctx = ssl.create_default_context()
    ctx.load_verify_locations(cafile=str(pem))
    return ctx


def descargar():
    CRUDO.mkdir(parents=True, exist_ok=True)
    destino = CRUDO / ARCHIVO
    ctx = contexto_tls()
    try:
        with urllib.request.urlopen(URL, timeout=60, context=ctx) as r:
            datos = r.read()
    except ssl.SSLCertVerificationError as e:
        raise SystemExit(
            f"La verificación TLS falló incluso con el intermedio: {e}. "
            f"Puede que ASECAM haya rotado su certificado; hay que traer el "
            f"nuevo intermedio de su extensión Authority Information Access.")
    if not datos.startswith(b"PK"):
        raise SystemExit(
            f"Lo descargado no es un XLSX (no empieza con PK). Primeros "
            f"bytes: {datos[:60]!r}")
    destino.write_bytes(datos)
    print(f"  {destino.relative_to(RAIZ)} · {len(datos):,} bytes · TLS verificado")


def clasificar(entidad):
    """(tipo, cve_mun). El municipio sólo se extrae de entidades municipales."""
    n = norm(entidad)
    for tipo, patron in PATRONES:
        m = re.match(patron, n)
        if m:
            # El nombre del municipio se cruza normalizado contra CVE_MUN,
            # nunca por posición ni por coincidencia parcial. Mismo criterio
            # que INE/INEGI en build_electoral.py.
            cve = CVE_MUN.get(norm(m.group(1)))
            return tipo, cve
    return "estatal", None


def leer(ruta):
    wb = openpyxl.load_workbook(ruta, read_only=True)
    filas = []
    for hoja in wb.sheetnames:
        ws = wb[hoja]
        cab = False
        for r in ws.iter_rows(values_only=True):
            v = [x for x in r]
            limpio = [x for x in v if x is not None]
            if not limpio:
                continue
            if str(limpio[0]).strip() == ENCABEZADO[0]:
                cab = True
                continue
            if not cab or len(limpio) < 2:
                continue
            entidad = str(v[1]).strip() if len(v) > 1 and v[1] else ""
            # El encabezado se repite dentro de la hoja y se colaría como si
            # fuera una entidad fiscalizada llamada "Entidad Fiscalizada".
            if not entidad or entidad == ENCABEZADO[1]:
                continue
            filas.append({
                "cuenta_publica": hoja,
                "entidad": entidad,
                "objeto": str(v[2]).strip() if len(v) > 2 and v[2] else None,
                "observacion": str(v[3]).strip() if len(v) > 3 and v[3] else None,
                "accion": str(v[4]).strip() if len(v) > 4 and v[4] else None,
            })
    return filas


def agregar(filas):
    por_mun, sin_mun, tipos = {}, [], {}
    for f in filas:
        tipo, cve = clasificar(f["entidad"])
        f["tipo_entidad"] = tipo
        f["cve_mun"] = cve
        tipos[tipo] = tipos.get(tipo, 0) + 1
        if cve:
            m = por_mun.setdefault(cve, {
                "cve_mun": cve, "nombre": NOMBRE_MUN[cve],
                "observaciones": 0, "por_tipo": {}, "por_cuenta": {},
                "entidades": set(),
            })
            m["observaciones"] += 1
            m["por_tipo"][tipo] = m["por_tipo"].get(tipo, 0) + 1
            m["por_cuenta"][f["cuenta_publica"]] = \
                m["por_cuenta"].get(f["cuenta_publica"], 0) + 1
            m["entidades"].add(f["entidad"])
        elif tipo != "estatal":
            # Municipal por su tipo pero sin municipio identificable: es un
            # fallo del cruce y hay que verlo, no perderlo.
            sin_mun.append(f["entidad"])

    for m in por_mun.values():
        m["entidades"] = sorted(m["entidades"])
    return por_mun, sorted(set(sin_mun)), tipos


def verificar(filas, por_mun, sin_mun, tipos):
    if not filas:
        raise SystemExit("No se leyó ninguna observación del archivo.")
    if sum(tipos.values()) != len(filas):
        raise SystemExit("La clasificación por tipo no cubre todas las filas.")
    if sin_mun:
        raise SystemExit(
            f"Entidades municipales cuyo municipio no se pudo identificar: "
            f"{sin_mun}. Revisa PATRONES o CVE_MUN antes de publicar: una "
            f"observación municipal que no se le asigna a nadie desaparece.")
    total_mun = sum(m["observaciones"] for m in por_mun.values())
    esperado = tipos.get("ayuntamiento", 0) + tipos.get("organismo_municipal", 0)
    if total_mun != esperado:
        raise SystemExit(
            f"Las observaciones municipales agregadas ({total_mun}) no cuadran "
            f"con las clasificadas como municipales ({esperado}).")
    for m in por_mun.values():
        if m["cve_mun"] not in NOMBRE_MUN:
            raise SystemExit(f"Clave de municipio desconocida: {m['cve_mun']}.")
        if sum(m["por_tipo"].values()) != m["observaciones"]:
            raise SystemExit(f"{m['nombre']}: el desglose por tipo no suma.")


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--descargar", action="store_true",
                    help="baja el archivo de ASECAM con TLS verificado")
    ap.add_argument("--autocomprobar", action="store_true")
    args = ap.parse_args()
    if args.autocomprobar:
        return autocomprobar()

    if args.descargar:
        print("Descargando de ASECAM…")
        descargar()

    ruta = CRUDO / ARCHIVO
    if not ruta.exists():
        raise SystemExit(f"Falta {ruta.relative_to(RAIZ)}. Corre con --descargar.")

    filas = leer(ruta)
    por_mun, sin_mun, tipos = agregar(filas)
    verificar(filas, por_mun, sin_mun, tipos)

    cuentas = sorted({f["cuenta_publica"] for f in filas})
    print(f"{len(filas)} observaciones · {len(cuentas)} cuentas públicas")
    print("  por tipo de entidad: " + " · ".join(f"{k} {v}" for k, v in tipos.items()))
    print(f"  municipios con observaciones: {len(por_mun)} de {len(NOMBRE_MUN)}")
    for m in sorted(por_mun.values(), key=lambda x: -x["observaciones"]):
        print(f"    {m['nombre']:14} {m['observaciones']:3}  "
              + " · ".join(f"{k} {v}" for k, v in sorted(m["por_tipo"].items())))
    sin_obs = [NOMBRE_MUN[c] for c in NOMBRE_MUN if c not in por_mun]
    print(f"  sin observaciones en estas cuentas: {sin_obs}")

    SALIDA.write_text(json.dumps({
        "metadata": {
            "procedencia": "dato",
            "fuente": FUENTE,
            "url": URL,
            "fecha_corte": FECHA_CORTE,
            "cuentas_publicas": cuentas,
            "que_es": "Observaciones de fiscalización externa sobre el "
                      "ejercicio de recursos públicos.",
            "no_es": "Ni una sanción ni una sentencia. Un pliego de "
                     "observaciones inicia un procedimiento; no lo resuelve. "
                     "El conteo no dice que haya habido un daño, y menos aún "
                     "quién lo causó.",
            "no_es_gasto_de_campana":
                "Es fiscalización sobre el municipio, no gasto de campaña. Son "
                "dos cosas distintas aunque las dos sean dinero: por eso esto "
                "alimenta el grafo de riesgos y no el módulo de finanzas.",
            "clasificacion":
                "Por TIPO de entidad, no por si el nombre menciona un "
                "municipio. 'Instituto Tecnológico Superior de Champotón' es "
                "estatal y está en Champotón: colgarle esa observación al "
                "municipio sería atribuirle algo que no es suyo.",
            "cruce": "El municipio se cruza por nombre normalizado contra la "
                     "misma tabla CVE_MUN que usa build_electoral.py.",
            "tls": {
                "problema": "asecam.gob.mx sirve su certificado sin el "
                            "intermedio: la verificación estándar falla con "
                            "'unable to get local issuer certificate'.",
                "solucion": "Se añade el intermedio —declarado por el propio "
                            "certificado en su extensión Authority Information "
                            "Access— al almacén de confianza. La verificación "
                            "sigue ACTIVA.",
                "lo_que_no_se_hizo": "Desactivar la verificación. Eso "
                                     "convertiría la descarga de una fuente "
                                     "oficial en la de cualquiera.",
                "intermedio": INTERMEDIO,
            },
            "vias_descartadas": {
                "datos_gob_mx": "Su API de CKAN devuelve 403 desde este "
                                "entorno. No responde.",
                "asf_federal": "asf.gob.mx publica 37 PDF y cero archivos "
                               "procesables; asfdatos.gob.mx es un sistema de "
                               "consulta web, no un catálogo de descarga. "
                               "BLOQUEADO por falta de fuente procesable.",
                "asecam_anos_previos": "Los 10 archivos anteriores tienen otro "
                                       "esquema (2022 trae una hoja 'Pliegos' "
                                       "con otras columnas). Fuera hasta "
                                       "verificar el mapeo.",
            },
            "conteo_por_tipo": tipos,
            "municipios_sin_observaciones": sin_obs,
            "generado_por": "backend/datalab/pipeline/build_auditoria.py",
        },
        "por_municipio": por_mun,
        "observaciones": filas,
    }, ensure_ascii=False), encoding="utf-8")
    print(f"Listo. {SALIDA.relative_to(RAIZ)}")


def autocomprobar():
    # La clasificación distingue tipo de entidad, que es lo delicado.
    casos = [
        ("Ayuntamiento de Calkiní", "ayuntamiento", "04001"),
        ("H. Ayuntamiento de Hecelchakán", "ayuntamiento", "04005"),
        ("Sistema Municipal de Agua Potable y Alcantarillado de Campeche",
         "organismo_municipal", "04002"),
        ("Sistema para el Desarrollo Integral de la Familia en el Municipio de Palizada",
         "organismo_municipal", "04007"),
        ("Sistema Municipal para el Desarrollo Integral de la Familia en Candelaria",
         "organismo_municipal", "04011"),
        # El que motiva la regla: estatal aunque nombre un municipio.
        ("Instituto Tecnológico Superior de Champotón", "estatal", None),
        ("Tribunal Electoral del Estado de Campeche", "estatal", None),
        ("Universidad Autónoma del Carmen", "estatal", None),
    ]
    for entidad, tipo_esp, cve_esp in casos:
        tipo, cve = clasificar(entidad)
        assert tipo == tipo_esp, f"{entidad}: {tipo} != {tipo_esp}"
        assert cve == cve_esp, f"{entidad}: {cve} != {cve_esp}"

    # El intermedio tiene que estar en el repositorio, o la descarga no puede
    # verificar y el pipeline debe detenerse.
    assert (CRUDO / INTERMEDIO).exists(), "falta el certificado intermedio"

    ruta = CRUDO / ARCHIVO
    if ruta.exists():
        filas = leer(ruta)
        por_mun, sin_mun, tipos = agregar(filas)
        verificar(filas, por_mun, sin_mun, tipos)
        # El encabezado repetido no puede colarse como entidad.
        assert not any(f["entidad"] == ENCABEZADO[1] for f in filas)
        print(f"OK · {len(filas)} observaciones · {len(por_mun)} municipios · "
              + " · ".join(f"{k} {v}" for k, v in tipos.items()))
    else:
        print("OK · clasificación verificada (falta el crudo: usa --descargar)")


if __name__ == "__main__":
    main()
